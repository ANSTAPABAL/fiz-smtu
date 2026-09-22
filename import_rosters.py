"""Import the supplied ISU roster exports into the local SQLite database.

The source spreadsheets stay outside the repository. Their student data is
stored only in the ignored local fkis.db file.
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

from server import connect, init_db


FILES = {
    "0901": 2,
    "0902": 3,
    "0903": 4,
    "0904": 5,
    "0905": 6,
    "0906": 7,
    "0907": 8,
    "0908": 9,
    "1001": 10,
    "1501": 11,
    "1502": 12,
    "1503": 13,
    "1507": 15,
    "2634": 16,
}

# The abbreviated filenames are mapped to the full groups and education rows
# shown in the university's official 2026/27 schedule matrix. The abbreviated
# key remains the private local DB key to avoid colliding with legacy demo rows.
FACULTY = "Факультет цифровых промышленных технологий"
GROUP_STRUCTURE = {
    **{f"090{i}": (f"12509-0{i}", FACULTY, "09.03.01 Информатика и вычислительная техника", 1) for i in range(1, 9)},
    "1001": ("12510-01", FACULTY, "10.03.01 Информационная безопасность", 1),
    "1501": ("12515-01", FACULTY, "15.03.01 Машиностроение", 1),
    "1502": ("12515-02", FACULTY, "15.03.01 Машиностроение", 1),
    "1503": ("12515-03", FACULTY, "15.03.01 Машиностроение", 1),
    "1507": ("12515-07", FACULTY, "15.03.06 Мехатроника и робототехника", 1),
    "2634": ("12526-34", FACULTY, "26.05.01 Проектирование и постройка кораблей, судов и объектов океанотехники", 1),
}


def main():
    source_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Downloads"
    init_db()
    imported = 0
    skipped = []
    con = connect()
    try:
        for group_code, file_number in FILES.items():
            path = source_dir / f"ИСУ СПбГМТУ в Санкт-Петербурге ({file_number}).xlsx"
            if not path.exists():
                skipped.append(path.name)
                continue
            display_code, faculty, direction, course = GROUP_STRUCTURE[group_code]
            con.execute(
                """INSERT INTO academic_groups(group_code,source,display_code,faculty,direction,course)
                   VALUES(?,?,?,?,?,?) ON CONFLICT(group_code) DO UPDATE SET
                   source=excluded.source,display_code=excluded.display_code,
                   faculty=excluded.faculty,direction=excluded.direction,course=excluded.course""",
                (group_code, "ИСУ + расписание СПбГМТУ", display_code, faculty, direction, course),
            )
            sheet = load_workbook(path, read_only=True, data_only=True).active
            # Row one is the export title and row two contains the six field names.
            for row in sheet.iter_rows(min_row=3, values_only=True):
                if not row or not row[1] or not row[5]:
                    continue
                position, isu_id, enrollment, study_form, _photo, name = row[:6]
                values = (str(name).strip(), group_code, str(enrollment or ""),
                          str(study_form or ""), int(position) if position else None)
                existing = con.execute("SELECT id FROM students WHERE isu_id=?", (str(isu_id),)).fetchone()
                if existing:
                    con.execute("UPDATE students SET name=?,group_code=?,enrollment_number=?,study_form=?,list_position=? WHERE id=?", (*values, existing[0]))
                else:
                    con.execute(
                        """INSERT INTO students
                           (name,health,attendance,theory,practice,group_code,isu_id,
                            enrollment_number,study_form,list_position)
                           VALUES(?,?,?,?,?,?,?,?,?,?)""",
                        (values[0], "не указана", 0, "не указано", "не указано",
                         values[1], str(isu_id), values[2], values[3], values[4]),
                    )
                imported += 1
        con.commit()
    finally:
        con.close()
    print(f"Обработано записей ИСУ: {imported}")
    if skipped:
        print("Не найдены файлы: " + ", ".join(skipped))
    print("Списки хранятся в локальной SQLite-базе fkis.db.")


if __name__ == "__main__":
    main()
