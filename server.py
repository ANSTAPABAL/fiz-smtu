import json
import sqlite3
import base64
import os
import hmac
import hashlib
import time
import re
from datetime import date
from io import BytesIO
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from urllib.parse import quote
from http.cookies import SimpleCookie
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DB = ROOT / "fkis.db"
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
APP_ENV = os.environ.get("APP_ENV", "development")

try:
    import psycopg
except ImportError:
    psycopg = None

if DATABASE_URL and psycopg is None:
    raise RuntimeError("Установите зависимости из requirements.txt для подключения к PostgreSQL.")

DB_ERRORS = (sqlite3.Error,) + ((psycopg.Error,) if psycopg else ())
APP_USERNAME = os.environ.get("APP_USERNAME", "fizruk")
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")
SESSION_SECRET = os.environ.get("SESSION_SECRET", "local-development-only")
AUTH_ENABLED = bool(APP_PASSWORD)
if APP_ENV == "production" and (not DATABASE_URL or not APP_PASSWORD or not os.environ.get("SESSION_SECRET")):
    raise RuntimeError("Для production задайте DATABASE_URL, APP_PASSWORD и SESSION_SECRET.")


class HybridRow(dict):
    """Psycopg row that preserves sqlite3.Row-style name and numeric access."""
    def __getitem__(self, key):
        if isinstance(key, int):
            return list(self.values())[key]
        return super().__getitem__(key)


def postgres_row_factory(cursor):
    columns = [column.name for column in (cursor.description or [])]
    return lambda values: HybridRow(zip(columns, values))


class PostgresConnection:
    def __init__(self, connection):
        self.connection = connection

    def execute(self, sql, args=()):
        sql = sql.replace("?", "%s")
        converted = tuple(
            date.fromisoformat(value) if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value) else value
            for value in args
        )
        return self.connection.execute(sql, converted)

    def commit(self): self.connection.commit()
    def rollback(self): self.connection.rollback()
    def close(self): self.connection.close()

DEFAULT_GROUP = "0905"
ACHIEVEMENT_CATEGORIES = ["Спортивное звание (разряд)", "ВФСК «ГТО»", "Сборная команда (секция)", "Спортивное мероприятие"]

def connect():
    if DATABASE_URL:
        return PostgresConnection(psycopg.connect(DATABASE_URL, row_factory=postgres_row_factory, connect_timeout=10))
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con

def init_db():
    if DATABASE_URL:
        con = connect()
        try:
            required = ("academic_groups", "students", "attendance_log", "achievements", "physical_tests")
            missing = [name for name in required if not con.execute("SELECT to_regclass(?) AS table_name", (f"public.{name}",)).fetchone()["table_name"]]
            if missing:
                raise RuntimeError("В удалённой базе отсутствуют таблицы: " + ", ".join(missing) + ". Выполните SQL из supabase/schema.sql.")
        finally:
            con.close()
        return
    con = connect()
    con.executescript("""
      CREATE TABLE IF NOT EXISTS students (id INTEGER PRIMARY KEY, name TEXT NOT NULL, health TEXT NOT NULL, attendance INTEGER NOT NULL DEFAULT 0, theory TEXT NOT NULL DEFAULT 'не указано', practice TEXT NOT NULL DEFAULT 'не указано', group_code TEXT NOT NULL DEFAULT '0905');
      CREATE TABLE IF NOT EXISTS attendance_log (id INTEGER PRIMARY KEY, student_id INTEGER NOT NULL REFERENCES students(id), lesson_date TEXT NOT NULL, topic TEXT NOT NULL, present INTEGER NOT NULL, UNIQUE(student_id, lesson_date));
      CREATE TABLE IF NOT EXISTS achievements (id INTEGER PRIMARY KEY, student_id INTEGER NOT NULL REFERENCES students(id), category TEXT NOT NULL, details TEXT NOT NULL, record_date TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'Подтверждено');
      CREATE TABLE IF NOT EXISTS physical_tests (id INTEGER PRIMARY KEY, student_id INTEGER NOT NULL REFERENCES students(id), record_date TEXT NOT NULL, exercise TEXT NOT NULL, result TEXT NOT NULL, UNIQUE(student_id,record_date,exercise));
      CREATE TABLE IF NOT EXISTS academic_groups (group_code TEXT PRIMARY KEY, source TEXT NOT NULL DEFAULT 'local');
      CREATE INDEX IF NOT EXISTS idx_attendance_student_date ON attendance_log(student_id, lesson_date);
      CREATE INDEX IF NOT EXISTS idx_achievements_student ON achievements(student_id);
    """)
    group_columns = [r[1] for r in con.execute("PRAGMA table_info(academic_groups)")]
    for name, declaration in {
        "display_code": "TEXT", "faculty": "TEXT", "direction": "TEXT", "course": "INTEGER"
    }.items():
        if name not in group_columns:
            con.execute(f"ALTER TABLE academic_groups ADD COLUMN {name} {declaration}")
    columns = [r[1] for r in con.execute("PRAGMA table_info(students)")]
    if "group_code" not in columns:
        con.execute("ALTER TABLE students ADD COLUMN group_code TEXT NOT NULL DEFAULT '12509-05'")
    for name, declaration in {
        "isu_id": "TEXT", "enrollment_number": "TEXT", "study_form": "TEXT",
        "list_position": "INTEGER", "birth_date": "TEXT", "theory_date": "TEXT"
    }.items():
        if name not in columns:
            con.execute(f"ALTER TABLE students ADD COLUMN {name} {declaration}")
    achievement_columns = [r[1] for r in con.execute("PRAGMA table_info(achievements)")]
    for name, declaration in {
        "sport_type": "TEXT", "distinction": "TEXT", "age_group": "TEXT",
        "order_basis": "TEXT", "participant_role": "TEXT", "event_result": "TEXT",
        "note": "TEXT", "document_name": "TEXT", "document_data": "BLOB",
        "team_name": "TEXT"
    }.items():
        if name not in achievement_columns:
            con.execute(f"ALTER TABLE achievements ADD COLUMN {name} {declaration}")
    con.execute("CREATE INDEX IF NOT EXISTS idx_students_group ON students(group_code)")
    con.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_students_isu_id ON students(isu_id) WHERE isu_id IS NOT NULL")
    con.commit(); con.close()

def rows(sql, args=()):
    con = connect(); result = [dict(r) for r in con.execute(sql,args).fetchall()]; con.close(); return result

def decode_pdf(value):
    if not value:
        return None
    document = base64.b64decode(value, validate=True)
    if len(document) > 10 * 1024 * 1024:
        raise ValueError("Файл не должен превышать 10 МБ.")
    if not document.startswith(b"%PDF-"):
        raise ValueError("Подтверждающий документ должен быть файлом PDF.")
    return document

class App(SimpleHTTPRequestHandler):
    def __init__(self,*args,**kwargs): super().__init__(*args, directory=str(ROOT), **kwargs)
    def send_json(self, value, status=200):
        raw=json.dumps(value,ensure_ascii=False,default=lambda item:item.isoformat() if hasattr(item,"isoformat") else str(item)).encode(); self.send_response(status); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Content-Length",str(len(raw))); self.end_headers(); self.wfile.write(raw)
    def body(self): return json.loads(self.rfile.read(int(self.headers.get("Content-Length",0))) or b"{}")
    def authenticated(self):
        if not AUTH_ENABLED:
            return True
        try:
            cookies = SimpleCookie(self.headers.get("Cookie", ""))
            token = cookies.get("fkis_session")
            if not token:
                return False
            payload, signature = token.value.rsplit(".", 1)
            expected = hmac.new(SESSION_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
            if not hmac.compare_digest(signature, expected):
                return False
            decoded = base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)).decode()
            username, expiry = decoded.rsplit("|", 1)
            return username == APP_USERNAME and int(expiry) > int(time.time())
        except (ValueError, TypeError, UnicodeDecodeError):
            return False
    def auth_required(self):
        if not AUTH_ENABLED or self.authenticated():
            return False
        self.send_json({"error":"Требуется войти в систему."},401)
        return True
    def issue_session(self):
        payload = base64.urlsafe_b64encode(f"{APP_USERNAME}|{int(time.time()) + 43200}".encode()).decode().rstrip("=")
        signature = hmac.new(SESSION_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
        return payload + "." + signature
    def do_GET(self):
        path=urlparse(self.path).path
        query=parse_qs(urlparse(self.path).query)
        if path == "/health":
            self.send_json({"status":"ok"}); return
        if path == "/api/session":
            self.send_json({"auth_enabled":AUTH_ENABLED,"authenticated":self.authenticated()}); return
        if path == "/login":
            self.path = "/login.html"
            return super().do_GET()
        if AUTH_ENABLED and not self.authenticated():
            public_assets = {"/styles.css", "/app.js", "/group-picker.js", "/achievement-filters.js", "/favicon.svg", "/smtu-logo.svg"}
            if path in public_assets:
                return super().do_GET()
            if path in ("/", "/index.html"):
                self.send_response(303); self.send_header("Location","/login"); self.end_headers(); return
            if path.startswith("/api/"):
                self.auth_required(); return
            self.send_json({"error":"Требуется войти в систему."},401); return
        group_code=query.get("group", [DEFAULT_GROUP])[0]
        if path == "/api/groups":
            self.send_json({"groups": rows("SELECT g.group_code,COALESCE(g.display_code,g.group_code) AS display_code,g.faculty,g.direction,g.course,COUNT(s.id) AS students FROM academic_groups g LEFT JOIN students s ON s.group_code=g.group_code GROUP BY g.group_code ORDER BY g.faculty,g.direction,g.course,g.display_code")}); return
        if not rows("SELECT group_code FROM academic_groups WHERE group_code=?",(group_code,)):
            group_code = DEFAULT_GROUP if rows("SELECT group_code FROM academic_groups WHERE group_code=?",(DEFAULT_GROUP,)) else ((rows("SELECT group_code FROM academic_groups ORDER BY group_code LIMIT 1") or [{"group_code":DEFAULT_GROUP}])[0]["group_code"])
        if path=="/api/export.xlsx":
            students=rows("SELECT s.*,CASE WHEN COUNT(a.id)>0 THEN ROUND(100.0*SUM(CASE WHEN a.present THEN 1 ELSE 0 END)/COUNT(a.id)) ELSE s.attendance END AS attendance_pct FROM students s LEFT JOIN attendance_log a ON a.student_id=s.id WHERE s.group_code=? GROUP BY s.id ORDER BY COALESCE(s.list_position,2147483647),s.name", (group_code,))
            achievements=rows("SELECT a.*,s.name,s.isu_id FROM achievements a JOIN students s ON s.id=a.student_id WHERE s.group_code=? ORDER BY s.list_position,a.id", (group_code,))
            tests=rows("SELECT p.*,s.name,s.isu_id,s.list_position FROM physical_tests p JOIN students s ON s.id=p.student_id WHERE s.group_code=? ORDER BY s.list_position,p.record_date,p.id", (group_code,))
            wb=Workbook(); overview=wb.active; overview.title="Карточка группы"
            roster=wb.create_sheet("Сводная ведомость"); fitness=wb.create_sheet("Физподготовка"); sport=wb.create_sheet("Достижения"); attendance=wb.create_sheet("Посещаемость")
            navy="455273"; blue="517CB3"; pale="EDF4FB"; line="D9E2F0"; white="FFFFFF"; ink="212529"
            thin=Side(style="thin",color=line)
            def title(sheet,label,columns):
                sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=columns)
                cell=sheet.cell(1,1,label); cell.font=Font(name="Fira Sans",size=16,bold=True,color=white); cell.fill=PatternFill("solid",fgColor=navy); cell.alignment=Alignment(horizontal="center",vertical="center")
                sheet.row_dimensions[1].height=32
                sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=columns)
                cell=sheet.cell(2,1,f"Группа {group_code} · сформировано {date.today().strftime('%d.%m.%Y')}"); cell.font=Font(name="Fira Sans",size=11,bold=True,color=navy); cell.fill=PatternFill("solid",fgColor=pale); cell.alignment=Alignment(horizontal="center")
                sheet.row_dimensions[2].height=23
                sheet.sheet_view.showGridLines=False
            def table(sheet,headers,data,widths):
                for col,label in enumerate(headers,1):
                    cell=sheet.cell(4,col,label); cell.font=Font(name="Fira Sans",bold=True,color=white); cell.fill=PatternFill("solid",fgColor=blue); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=Border(bottom=thin)
                sheet.row_dimensions[4].height=34
                for row_index,values in enumerate(data,5):
                    for col,value in enumerate(values,1):
                        cell=sheet.cell(row_index,col,value if value is not None else "—"); cell.font=Font(name="Fira Sans",size=10,color=ink); cell.alignment=Alignment(vertical="center",wrap_text=True,horizontal="center" if col==1 else "left"); cell.border=Border(bottom=thin)
                        if row_index%2: cell.fill=PatternFill("solid",fgColor="F8FAFC")
                    sheet.row_dimensions[row_index].height=28
                if data: sheet.auto_filter.ref=f"A4:{get_column_letter(len(headers))}{len(data)+4}"
                sheet.freeze_panes="A5"
                for col,width in enumerate(widths,1): sheet.column_dimensions[get_column_letter(col)].width=width
            total=len(students); present_logs=sum(int(row["present"] or 0) for row in rows("SELECT present FROM attendance_log WHERE student_id IN (SELECT id FROM students WHERE group_code=?)",(group_code,)))
            lesson_logs=int(rows("SELECT COUNT(*) AS n FROM attendance_log WHERE student_id IN (SELECT id FROM students WHERE group_code=?)",(group_code,))[0]["n"])
            ranks=sum(1 for a in achievements if "разряд" in a["category"].lower() or "звание" in a["category"].lower())
            gto=sum(1 for a in achievements if "гто" in a["category"].lower())
            teams=sum(1 for a in achievements if ("сборная" in a["category"].lower() or "секция" in a["category"].lower()) and a["status"]!="Исключен")
            events=sum(1 for a in achievements if "мероприят" in a["category"].lower())
            title(overview,"СПбГМТУ · Учет физической культуры и спорта",4)
            summary_rows=[["Состав","Студентов в группе",total],["Здоровье","Основная / подготовительная / специальная",f"{sum(1 for s in students if s['health']=='основная')} / {sum(1 for s in students if s['health']=='подготовительная')} / {sum(1 for s in students if s['health']=='специальная')}"],["Здоровье","Не указано",sum(1 for s in students if s['health']=='не указана')],["ТиМ ФКиС","Зачет / не зачет / не указано",f"{sum(1 for s in students if s['theory']=='зачет')} / {sum(1 for s in students if s['theory']=='не зачет')} / {sum(1 for s in students if s['theory']=='не указано')}"],["Практика ФКиС","Зачет / не зачет / не указано",f"{sum(1 for s in students if s['practice']=='зачет')} / {sum(1 for s in students if s['practice']=='не зачет')} / {sum(1 for s in students if s['practice']=='не указано')}"],["Посещаемость","Отметки присутствия",f"{present_logs} из {lesson_logs}"],["Посещаемость","Доля присутствия",f"{round(present_logs/lesson_logs*100)}%" if lesson_logs else "нет отметок"],["Физическая подготовленность","Записей испытаний",len(tests)],["Спортивное звание / разряд","Записей",ranks],["ВФСК «ГТО»","Записей",gto],["Сборная / спортивная секция","Действующих записей",teams],["Спортивное мероприятие","Записей участника / волонтера",events]]
            table(overview,["Раздел","Показатель","Значение"],summary_rows,[30,45,28])
            roster_rows=[]
            for index,s in enumerate(students,1):
                records=[a for a in achievements if a["student_id"]==s["id"]]
                team=next((a for a in records if "сборная" in a["category"].lower() or "секция" in a["category"].lower()),None)
                gto_row=next((a for a in records if "гто" in a["category"].lower()),None)
                rank=next((a for a in records if "разряд" in a["category"].lower() or "звание" in a["category"].lower()),None)
                event_rows=[a for a in records if "мероприят" in a["category"].lower()]
                roster_rows.append([index,s["isu_id"],s["enrollment_number"],s["name"],s["birth_date"],s["health"],s["theory"],s["theory_date"],s["practice"],s["attendance_pct"]/100,team["team_name"] if team else "—",gto_row["distinction"] if gto_row else "—",rank["distinction"] if rank else "—","; ".join(a["details"] or a["event_result"] or "" for a in event_rows) or "—",sum(1 for a in event_rows if a["participant_role"]=="Участник"),sum(1 for a in event_rows if a["participant_role"]=="Волонтер")])
            title(roster,"Сводная карточка учебной группы",16)
            table(roster,["№","ISU ID","Номер ЗК","ФИО","Дата рождения","Группа здоровья","ТиМ ФКиС","Дата ТиМ","Практика ФКиС","Посещаемость","Сборная команда","ВФСК «ГТО»","Спортивный разряд","Спортивное достижение","Участник","Волонтер"],roster_rows,[7,14,16,34,16,20,18,14,18,16,25,22,24,32,12,12])
            fitness_rows=[[next((i for i,student in enumerate(students,1) if student["id"]==test["student_id"]),"—"),test["isu_id"],test["name"],test["exercise"],test["result"],test["record_date"]] for test in tests]
            title(fitness,"Физическая подготовленность",6)
            table(fitness,["№","ISU ID","ФИО","Физическое упражнение","Результат","Дата выполнения"],fitness_rows,[7,14,34,36,25,18])
            sport_rows=[[next((i for i,student in enumerate(students,1) if student["id"]==item["student_id"]),"—"),item["isu_id"],item["name"],item["category"],item["sport_type"],item["team_name"],item["distinction"],item["age_group"],item["participant_role"],item["details"],item["event_result"],item["record_date"],item["order_basis"],item["document_name"],item["status"],item["note"]] for item in achievements]
            title(sport,"Спортивные достижения студентов",16)
            table(sport,["№","ISU ID","ФИО","Категория","Вид спорта / секция","Сборная","Звание / знак","Возрастная группа","Роль","Мероприятие","Результат / место","Дата","Приказ / основание","Подтверждающий документ","Состояние","Примечание"],sport_rows,[7,14,32,27,25,24,26,27,14,30,22,16,24,32,18,32])
            attendance_rows=rows("SELECT s.list_position,s.isu_id,s.name,s.health,a.lesson_date,a.topic,a.present FROM attendance_log a JOIN students s ON s.id=a.student_id WHERE s.group_code=? ORDER BY a.lesson_date DESC,s.list_position",(group_code,))
            attendance_data=[[item["list_position"],item["isu_id"],item["name"],item["health"],item["lesson_date"],item["topic"],"Присутствовал" if item["present"] else "Отсутствовал"] for item in attendance_rows]
            title(attendance,"Журнал посещаемости",7)
            table(attendance,["№","ISU ID","ФИО","Группа здоровья","Дата занятия","Дисциплина","Статус"],attendance_data,[7,14,34,21,17,30,20])
            stream=BytesIO(); wb.save(stream); raw=stream.getvalue()
            self.send_response(200); self.send_header("Content-Type","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); self.send_header("Content-Disposition",f"attachment; filename=fkis-{group_code}.xlsx"); self.send_header("Content-Length",str(len(raw))); self.end_headers(); self.wfile.write(raw); return
        if path=="/api/bootstrap":
            students=rows("SELECT students.*,CASE WHEN COUNT(a.id)>0 THEN ROUND(100.0*SUM(CASE WHEN a.present THEN 1 ELSE 0 END)/COUNT(a.id)) ELSE students.attendance END AS computed_attendance,COUNT(a.id) AS attendance_records_count FROM students LEFT JOIN attendance_log a ON a.student_id=students.id WHERE group_code=? GROUP BY students.id ORDER BY COALESCE(list_position,2147483647),name", (group_code,))
            achievements=rows("SELECT a.id,a.student_id,s.name,a.category,a.details,a.record_date,a.status,a.sport_type,a.distinction,a.age_group,a.order_basis,a.participant_role,a.event_result,a.note,a.document_name,a.team_name FROM achievements a JOIN students s ON s.id=a.student_id WHERE s.group_code=? ORDER BY a.id DESC", (group_code,))
            physical_tests=rows("SELECT p.*,s.name FROM physical_tests p JOIN students s ON s.id=p.student_id WHERE s.group_code=? ORDER BY p.record_date DESC,p.id DESC", (group_code,))
            attendance_records=rows("SELECT student_id,present,topic FROM attendance_log WHERE lesson_date=? AND student_id IN (SELECT id FROM students WHERE group_code=?)", (query.get("date", [date.today().isoformat()])[0], group_code))
            attendance_trend=rows("SELECT lesson_date,COUNT(*) AS total,SUM(CASE WHEN present THEN 1 ELSE 0 END) AS present FROM attendance_log WHERE student_id IN (SELECT id FROM students WHERE group_code=?) GROUP BY lesson_date ORDER BY lesson_date DESC LIMIT 7", (group_code,))
            summary = {
                "health": rows("SELECT health AS label,COUNT(*) AS count FROM students WHERE group_code=? GROUP BY health ORDER BY health", (group_code,)),
                "theory": rows("SELECT theory AS label,COUNT(*) AS count FROM students WHERE group_code=? GROUP BY theory ORDER BY theory", (group_code,)),
                "practice": rows("SELECT practice AS label,COUNT(*) AS count FROM students WHERE group_code=? GROUP BY practice ORDER BY practice", (group_code,)),
                "achievements": rows("SELECT category AS label,COUNT(*) AS count FROM achievements a JOIN students s ON s.id=a.student_id WHERE s.group_code=? GROUP BY category ORDER BY category", (group_code,)),
                "attendance": rows("SELECT COUNT(*) AS lessons,SUM(CASE WHEN present THEN 1 ELSE 0 END) AS present FROM attendance_log WHERE student_id IN (SELECT id FROM students WHERE group_code=?)", (group_code,))[0]
            }
            self.send_json({"students":students,"achievements":achievements,"physical_tests":physical_tests,"attendance_records":attendance_records,"attendance_trend":attendance_trend,"summary":summary,"group":group_code,"today":date.today().isoformat()}); return
        if path.startswith("/api/achievement-document/"):
            achievement_id=int(path.rsplit("/",1)[1]); con=connect(); item=con.execute("SELECT document_name,document_data FROM achievements WHERE id=?",(achievement_id,)).fetchone(); con.close()
            if not item or not item[1]: self.send_error(404); return
            raw=bytes(item[1]); name=quote(item[0] or "document.pdf")
            self.send_response(200); self.send_header("Content-Type","application/pdf"); self.send_header("Content-Disposition",f"attachment; filename*=UTF-8''{name}"); self.send_header("Content-Length",str(len(raw))); self.end_headers(); self.wfile.write(raw); return
        return super().do_GET()
    def do_PUT(self):
        if self.auth_required(): return
        path=urlparse(self.path).path; data=self.body(); con=connect()
        try:
            if path.startswith("/api/students/"):
                student_id=int(path.rsplit("/",1)[1])
                con.execute("UPDATE students SET name=?,health=?,attendance=?,theory=?,practice=?,birth_date=?,theory_date=? WHERE id=?",(data["name"].strip(),data["health"],int(data["attendance"]),data["theory"],data["practice"],data.get("birth_date") or None,data.get("theory_date") or None,student_id))
                con.commit(); self.send_json({"ok":True}); return
            if path.startswith("/api/achievements/"):
                achievement_id=int(path.rsplit("/",1)[1])
                document=decode_pdf(data.get("document_base64"))
                con.execute("UPDATE achievements SET student_id=?,category=?,details=?,record_date=?,status=?,sport_type=?,distinction=?,age_group=?,order_basis=?,participant_role=?,event_result=?,note=?,document_name=CASE WHEN ? THEN NULL WHEN ? IS NOT NULL THEN ? ELSE document_name END,document_data=CASE WHEN ? THEN NULL WHEN ? IS NOT NULL THEN ? ELSE document_data END,team_name=? WHERE id=?",(int(data["student_id"]),data["category"],data.get("details","").strip(),data["record_date"],data["status"],data.get("sport_type"),data.get("distinction"),data.get("age_group"),data.get("order_basis"),data.get("participant_role"),data.get("event_result"),data.get("note"),bool(data.get("remove_document")),document,data.get("document_name"),bool(data.get("remove_document")),document,document,data.get("team_name"),achievement_id))
                con.commit(); self.send_json({"ok":True}); return
            self.send_json({"error":"Неизвестный запрос"},404)
        except (KeyError, ValueError, *DB_ERRORS, json.JSONDecodeError, base64.binascii.Error) as e: con.rollback(); self.send_json({"error":str(e)},400)
        finally: con.close()
    def do_POST(self):
        path=urlparse(self.path).path; data=self.body()
        if path=="/api/login":
            username=str(data.get("username", "")); password=str(data.get("password", ""))
            if not AUTH_ENABLED or not hmac.compare_digest(username, APP_USERNAME) or not hmac.compare_digest(password, APP_PASSWORD):
                self.send_json({"error":"Неверный логин или пароль."},401); return
            token=self.issue_session(); secure=self.headers.get("X-Forwarded-Proto", "").lower()=="https" or bool(os.environ.get("RENDER"))
            self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Set-Cookie", f"fkis_session={token}; Path=/; HttpOnly; SameSite=Strict; Max-Age=43200" + ("; Secure" if secure else "")); self.end_headers(); self.wfile.write(b'{"ok":true}'); return
        if path=="/api/logout":
            self.send_response(200); self.send_header("Content-Type","application/json; charset=utf-8"); self.send_header("Set-Cookie","fkis_session=; Path=/; HttpOnly; SameSite=Strict; Max-Age=0"); self.end_headers(); self.wfile.write(b'{"ok":true}'); return
        if self.auth_required(): return
        con=connect()
        try:
            if path=="/api/attendance":
                for item in data.get("items",[]): con.execute("INSERT INTO attendance_log(student_id,lesson_date,topic,present) VALUES(?,?,?,?) ON CONFLICT(student_id,lesson_date) DO UPDATE SET topic=excluded.topic,present=excluded.present",(item["id"],data.get("lesson_date","2026-09-16"),data.get("topic","Общая физическая подготовка"),bool(item["present"])))
                con.commit(); self.send_json({"ok":True}); return
            if path=="/api/physical-tests":
                con.execute("INSERT INTO physical_tests(student_id,record_date,exercise,result) VALUES(?,?,?,?) ON CONFLICT(student_id,record_date,exercise) DO UPDATE SET result=excluded.result",(int(data["student_id"]),data["record_date"],data["exercise"].strip(),data["result"].strip()))
                con.commit(); self.send_json({"ok":True}); return
            if path=="/api/achievements":
                if data["category"] not in ACHIEVEMENT_CATEGORIES: raise ValueError("Выберите раздел спортивного учета.")
                document=decode_pdf(data.get("document_base64"))
                con.execute("INSERT INTO achievements(student_id,category,details,record_date,status,sport_type,distinction,age_group,order_basis,participant_role,event_result,note,document_name,document_data,team_name) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",(int(data["student_id"]),data["category"],data.get("details","").strip(),data.get("record_date") or date.today().strftime("%d.%m.%Y"),data.get("status","Подтверждено"),data.get("sport_type"),data.get("distinction"),data.get("age_group"),data.get("order_basis"),data.get("participant_role"),data.get("event_result"),data.get("note"),data.get("document_name"),document,data.get("team_name"))); con.commit(); self.send_json({"ok":True},201); return
            if path=="/api/students":
                group_code=data.get("group_code",DEFAULT_GROUP)
                if not con.execute("SELECT 1 FROM academic_groups WHERE group_code=?",(group_code,)).fetchone(): raise ValueError("Выберите группу из списка ИСУ.")
                cur=con.execute("INSERT INTO students(name,health,attendance,theory,practice,group_code) VALUES(?,?,?,?,?,?) RETURNING id",(data["name"],data.get("health","не указана"),0,"не указано","не указано",group_code)); student_id=cur.fetchone()["id"]; con.commit(); self.send_json({"id":student_id},201); return
            self.send_json({"error":"Неизвестный запрос"},404)
        except (KeyError, ValueError, *DB_ERRORS, json.JSONDecodeError, base64.binascii.Error) as e: con.rollback(); self.send_json({"error":str(e)},400)
        finally: con.close()

if __name__=="__main__":
    init_db()
    port=int(os.environ.get("PORT", "4174"))
    print(f"http://127.0.0.1:{port}")
    bind_host="0.0.0.0" if os.environ.get("PORT") or DATABASE_URL else "127.0.0.1"
    ThreadingHTTPServer((bind_host,port),App).serve_forever()
